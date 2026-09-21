import psycopg2
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

DATABASE_URL = "postgresql://neondb_owner:npg_Rl0zv1crIkTY@ep-purple-heart-axjsakzf-pooler.c-4.us-east-2.aws.neon.tech/neondb?sslmode=require"

def export_photocards():
    conn = psycopg2.connect(DATABASE_URL)
    cur = conn.cursor()
    
    # 1. Total statistics
    cur.execute("SELECT COUNT(*) FROM candidate_roster")
    total_candidates = cur.fetchone()[0]
    
    cur.execute("""
        SELECT COUNT(*) 
        FROM candidate_roster 
        WHERE registration_status IN ('registered', 'tested') OR registered_at IS NOT NULL
    """)
    total_registered = cur.fetchone()[0]
    
    cur.execute("""
        SELECT 
            psn, 
            name, 
            amended_name, 
            mda, 
            proposed_rank, 
            proposed_gl, 
            exam_code,
            phone, 
            email, 
            registration_status, 
            registered_at
        FROM candidate_roster
        WHERE registration_status IN ('registered', 'tested') OR registered_at IS NOT NULL
        ORDER BY registered_at DESC
    """)
    rows = cur.fetchall()
    conn.close()

    print(f"Total Candidates on Roster: {total_candidates}")
    print(f"Total Candidates Verified & Photocard Generated: {total_registered} ({total_registered/total_candidates*100:.2f}%)")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Photocards Generated"
    ws.views.sheetView[0].showGridLines = True

    primary_green = "004D40"
    light_green = "E0F2F1"
    border_gray = "CCCCCC"
    thin_border = Border(
        left=Side(style='thin', color=border_gray),
        right=Side(style='thin', color=border_gray),
        top=Side(style='thin', color=border_gray),
        bottom=Side(style='thin', color=border_gray)
    )

    # Title Block
    ws.merge_cells("A1:K1")
    ws["A1"] = "KWARA STATE CIVIL SERVICE COMMISSION"
    ws["A1"].font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=primary_green, end_color=primary_green, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:K2")
    ws["A2"] = "2026 Promotion CBT Examination - Verified Candidates & Photocard Generation Log"
    ws["A2"].font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    ws["A2"].fill = PatternFill(start_color="00796B", end_color="00796B", fill_type="solid")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    ws.merge_cells("A3:K3")
    ws["A3"] = f"Report Generated: {datetime.now().strftime('%d-%b-%Y %I:%M %p')} | Total Photocard Generated: {total_registered} of {total_candidates} ({total_registered/total_candidates*100:.2f}%)"
    ws["A3"].font = Font(name="Arial", size=10, italic=True, color="333333")
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 20

    headers = [
        "S/N", "PSN", "Roster Name", "Verified / Amended Name", 
        "MDA", "Proposed Rank", "GL", "Exam Code", 
        "Phone Number", "Email Address", "Photocard Generated At"
    ]

    header_row = 5
    ws.row_dimensions[header_row].height = 25
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=primary_green, end_color=primary_green, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for idx, r in enumerate(rows, 1):
        curr_row = header_row + idx
        ws.row_dimensions[curr_row].height = 20
        reg_time_str = r[10].strftime("%d-%b-%Y %I:%M:%S %p") if r[10] else "N/A"
        row_vals = [
            idx,
            r[0],
            r[1],
            r[2] or r[1],
            r[3],
            r[4],
            r[5],
            r[6],
            r[7] or "N/A",
            r[8] or "N/A",
            reg_time_str
        ]
        is_even = (idx % 2 == 0)
        row_fill = PatternFill(start_color=light_green if is_even else "FFFFFF", fill_type="solid")

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.font = Font(name="Arial", size=9)
            cell.fill = row_fill
            cell.border = thin_border
            if c_idx in [1, 2, 7, 8, 11]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:
                continue
            v = str(cell.value or "")
            if len(v) > max_len:
                max_len = len(v)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    filename = f"Kwara_CSC_2026_Photocard_Generated_Candidates_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    wb.save(filename)
    print(f"Excel report saved successfully to: {filename}")

if __name__ == "__main__":
    export_photocards()
