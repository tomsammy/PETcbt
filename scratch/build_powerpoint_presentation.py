import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    # Official Color Palette
    C_PRIMARY_DARK = RGBColor(0, 77, 64)      # #004D40 Deep Emerald
    C_PRIMARY_TEAL = RGBColor(0, 121, 107)    # #00796B
    C_GOLD = RGBColor(217, 119, 6)            # #D97706
    C_RED = RGBColor(220, 38, 38)             # #DC2626
    C_GREEN = RGBColor(16, 185, 129)          # #10B981
    C_DARK = RGBColor(30, 41, 59)             # #1E293B
    C_MUTED = RGBColor(100, 116, 139)         # #64748B
    C_BG_LIGHT = RGBColor(248, 250, 252)      # #F8FAFC
    C_WHITE = RGBColor(255, 255, 255)
    C_BORDER = RGBColor(226, 232, 240)        # #E2E8F0

    logo_path = os.path.abspath("kwara_cbt_app/static/images/logo.png")
    screenshot_dir = os.path.abspath("scratch/orientation_screenshots")

    def add_header_banner(slide, category_text, stage_title, slide_num_str):
        # Top banner background
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = C_PRIMARY_DARK
        top_bar.line.color.rgb = C_PRIMARY_DARK

        # Logo on left
        if os.path.exists(logo_path):
            slide.shapes.add_picture(logo_path, Inches(0.4), Inches(0.12), height=Inches(0.9))

        # Title text box
        tx_box = slide.shapes.add_textbox(Inches(1.45), Inches(0.12), Inches(9.8), Inches(0.95))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = category_text.upper()
        p1.font.name = "Arial"
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(167, 243, 208) # Mint light

        p2 = tf.add_paragraph()
        p2.text = stage_title
        p2.font.name = "Arial"
        p2.font.size = Pt(17)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE

        # Slide pill badge on right
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.4), Inches(0.35), Inches(1.5), Inches(0.45))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_PRIMARY_TEAL
        badge.line.color.rgb = RGBColor(167, 243, 208)
        bf = badge.text_frame
        bp = bf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        bp.text = slide_num_str
        bp.font.name = "Arial"
        bp.font.size = Pt(11)
        bp.font.bold = True
        bp.font.color.rgb = C_WHITE

    def add_footer(slide):
        footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.08), Inches(12.133), Inches(0.35))
        ftf = footer_box.text_frame
        p = ftf.paragraphs[0]
        p.text = "Kwara State Civil Service Commission (CSC) • 2026 Promotion Evaluation CBT Examination Guide"
        p.font.name = "Arial"
        p.font.size = Pt(9)
        p.font.color.rgb = C_MUTED

    def set_speaker_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    # Background full color
    bg = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_PRIMARY_DARK
    bg.line.color.rgb = C_PRIMARY_DARK

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.2), Inches(13.333), Inches(0.3))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_GOLD
    bar.line.color.rgb = C_GOLD

    # Logo
    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(5.9), Inches(0.9), height=Inches(1.5))

    # Title Box
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(11.333), Inches(3.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "KWARA STATE CIVIL SERVICE COMMISSION"
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(167, 243, 208)

    p = tf1.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "2026 Promotion Evaluation CBT Examination"
    p.font.name = "Arial"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    p = tf1.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "Candidate Step-by-Step Orientation & System Interface Walkthrough"
    p.font.name = "Arial"
    p.font.size = Pt(17)
    p.font.color.rgb = RGBColor(254, 243, 199) # Pale gold

    # Sub-card at bottom
    sub_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.5), Inches(5.6), Inches(8.333), Inches(0.9))
    sub_card.fill.solid()
    sub_card.fill.fore_color.rgb = C_PRIMARY_TEAL
    sub_card.line.color.rgb = RGBColor(167, 243, 208)
    stf = sub_card.text_frame
    sp = stf.paragraphs[0]
    sp.alignment = PP_ALIGN.CENTER
    sp.text = "Official Briefing for All Civil Service Promotion Candidates • Batch 2026"
    sp.font.name = "Arial"
    sp.font.size = Pt(13)
    sp.font.bold = True
    sp.font.color.rgb = C_WHITE

    set_speaker_notes(s1, 
        "FACILITATOR SCRIPT:\n"
        "'Good day, distinguished officers and candidates of the Kwara State Civil Service. "
        "Welcome to today's orientation session for the 2026 Promotion Evaluation Computer Based Test (CBT). "
        "The Civil Service Commission has structured this digital assessment to ensure merit, transparency, and seamless execution. "
        "In the next few minutes, we will walk you through every single screen and button you will encounter, "
        "the anti-cheating regulations, and what to expect on your test workstation. "
        "Please pay close attention to every detail.'"
    )

    # =========================================================================
    # SLIDE 2: EXAMINATION GROUND RULES & STRUCTURE
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header_banner(s2, "EXAMINATION OVERVIEW", "Core Ground Rules & Evaluation Structure", "Overview")
    add_footer(s2)

    # 4 Key Stat Cards
    cards_data = [
        ("⏱️", "20 MINUTES", "STRICT TIMING", "Each candidate has exactly 20 minutes to complete the evaluation. A live countdown clock is always visible on screen."),
        ("📝", "40 QUESTIONS", "MULTIPLE CHOICE", "Each paper consists of exactly 40 objective questions (1 mark each). All 40 questions should be attempted."),
        ("🔒", "ZERO SCORE ON SCREEN", "CONFIDENTIALITY", "In accordance with Civil Service rules, marks are NEVER displayed on workstation screens upon submission."),
        ("🛡️", "3-STRIKE POLICY", "PROCTOR SECURITY", "Switching apps, minimizing the browser, or taking screenshots triggers automated violation strikes.")
    ]

    for idx, (icon, val, lbl, desc) in enumerate(cards_data):
        x = Inches(0.6 + idx * 3.05)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.45), Inches(2.9), Inches(3.6))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.2)
        ctf.margin_top = Inches(0.25)

        p = ctf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = icon
        p.font.size = Pt(36)

        p = ctf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.text = val
        p.font.name = "Arial"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_PRIMARY_DARK

        p = ctf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.text = lbl
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_GOLD

        p = ctf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.text = desc
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK

    # Bottom Callout Box
    alert_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.3), Inches(12.133), Inches(1.55))
    alert_box.fill.solid()
    alert_box.fill.fore_color.rgb = RGBColor(254, 242, 242) # Soft red
    alert_box.line.color.rgb = RGBColor(254, 202, 202)
    alert_box.line.width = Pt(1.5)

    atf = alert_box.text_frame
    atf.word_wrap = True
    atf.margin_left = atf.margin_right = Inches(0.3)
    atf.margin_top = Inches(0.18)

    p = atf.paragraphs[0]
    p.text = "⚠️ MANDATORY REQUIREMENT BEFORE ENTERING THE HALL"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_RED

    p = atf.add_paragraph()
    p.text = (
        "• Bring your printed Photocard Registration Slip showing your Public Service Number (PSN) and Code 1.\n"
        "• Ensure you know your assigned Examination Batch and scheduled time slot.\n"
        "• If using a smartphone, ensure your battery is at least 80% charged and turn off background notifications."
    )
    p.font.name = "Arial"
    p.font.size = Pt(10.5)
    p.font.color.rgb = RGBColor(127, 29, 29)

    set_speaker_notes(s2,
        "FACILITATOR SCRIPT:\n"
        "'Before we examine the screens, remember these four pillars:\n"
        "1. Timing: You have strictly 20 minutes for 40 questions. That is an average of 30 seconds per question.\n"
        "2. Structure: 40 questions, objective multiple choice. Answer every single question.\n"
        "3. Result Policy: Do not expect a score on the computer when you finish. Scores are confidential and delivered to your MDA.\n"
        "4. Anti-cheating: The system detects tab switches and background applications. Three strikes terminate your exam.'"
    )

    # =========================================================================
    # SLIDE 3: STAGE 1 - CANDIDATE PORTAL LOGIN
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header_banner(s3, "STAGE 1 OF 7", "Candidate Portal Login & Accreditation", "Stage 1")
    add_footer(s3)

    # Screenshot on Left
    s1_img = os.path.join(screenshot_dir, "stage1_login.png")
    if os.path.exists(s1_img):
        s3.shapes.add_picture(s1_img, Inches(0.6), Inches(1.4), width=Inches(7.2))

    # Right side Instruction Panel
    panel = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.4), Inches(4.733), Inches(5.45))
    panel.fill.solid()
    panel.fill.fore_color.rgb = C_WHITE
    panel.line.color.rgb = C_BORDER
    panel.line.width = Pt(1.5)

    ptf = panel.text_frame
    ptf.word_wrap = True
    ptf.margin_left = ptf.margin_right = Inches(0.25)
    ptf.margin_top = Inches(0.25)

    p = ptf.paragraphs[0]
    p.text = "🔑 HOW TO LOG IN"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY_DARK

    steps = [
        ("Step 1: Check Portal Status", "Verify that the top status bar displays 'ACTIVE & ACCREDITATION OPEN'."),
        ("Step 2: Enter Public Service Number (PSN)", "Type your official 6-digit Kwara State civil service number (e.g. 123456)."),
        ("Step 3: Enter Official Code 1", "Enter the exact Examination Code 1 printed on your Photocard Slip (e.g. B-10294). This code is case-sensitive."),
        ("Step 4: Click 'Verify & Access'", "Click the dark green button to authenticate your candidate record.")
    ]

    for title, desc in steps:
        p = ptf.add_paragraph()
        p.text = f"• {title}: "
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_DARK
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_MUTED

    # Pro Tip box inside panel
    p = ptf.add_paragraph()
    p.text = "\n💡 PRO TIP FOR CANDIDATES:"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_GOLD

    p = ptf.add_paragraph()
    p.text = "Do not share your Code 1 with anyone. Each code can only be activated once for your specific examination paper."
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_DARK

    set_speaker_notes(s3,
        "FACILITATOR SCRIPT:\n"
        "'Here on Slide 3 is the exact login screen you will see on your computer or mobile screen.\n"
        "Look at your Photocard Slip. In the top right corner, you have your 6-digit PSN and your Code 1.\n"
        "Enter your PSN in the first box. Enter your Code 1 in the second box. Make sure you include any letters and hyphens.\n"
        "Then click 'Verify & Access Examination'. Do not click multiple times—wait 2 seconds for the accreditation window to load.'"
    )

    # =========================================================================
    # SLIDE 4: STAGE 2 - BIO-DATA & SLIP VERIFICATION
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header_banner(s4, "STAGE 2 OF 7", "Candidate Bio-Data & Paper Allocation Verification", "Stage 2")
    add_footer(s4)

    s2_img = os.path.join(screenshot_dir, "stage2_verification.png")
    if os.path.exists(s2_img):
        s4.shapes.add_picture(s2_img, Inches(0.6), Inches(1.4), width=Inches(7.2))

    panel = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.4), Inches(4.733), Inches(5.45))
    panel.fill.solid()
    panel.fill.fore_color.rgb = C_WHITE
    panel.line.color.rgb = C_BORDER
    panel.line.width = Pt(1.5)

    ptf = panel.text_frame
    ptf.word_wrap = True
    ptf.margin_left = ptf.margin_right = Inches(0.25)
    ptf.margin_top = Inches(0.25)

    p = ptf.paragraphs[0]
    p.text = "📋 VERIFY BEFORE YOU BEGIN"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY_DARK

    v_steps = [
        ("Full Name & PSN", "Ensure the name displayed matches your civil service record exactly."),
        ("Ministry / Department", "Verify your current posting/MDA is accurately reflected."),
        ("Grade Level & Cadre", "Confirm your proposed promotional rank and cadre."),
        ("Allocated Paper Code", "Check the paper code (e.g. MARD-GL12-B). This determines the questions presented."),
        ("If Everything is Correct", "Click the dark green 'Confirm & Begin Examination' button to launch the exam clock."),
        ("If You Notice an Error", "Click 'Report Discrepancy' immediately and alert the hall supervisor BEFORE starting.")
    ]

    for title, desc in v_steps:
        p = ptf.add_paragraph()
        p.text = f"✔ {title}: "
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_DARK
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_MUTED

    p = ptf.add_paragraph()
    p.text = "\n⚠️ CRITICAL WARNING:"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_RED

    p = ptf.add_paragraph()
    p.text = "Once you click 'Confirm & Begin Examination', your 20-minute timer begins counting down and CANNOT be paused or reset!"
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_RED

    set_speaker_notes(s4,
        "FACILITATOR SCRIPT:\n"
        "'When you authenticate, this verification pop-up appears.\n"
        "Do NOT rush to click 'Confirm & Begin'. Take 30 seconds to read every field.\n"
        "Is that your name? Is that your PSN? Is that your Ministry? Is that your allocated paper?\n"
        "If there is any discrepancy whatsoever, raise your hand immediately and inform the supervisor.\n"
        "Once you click 'Confirm & Begin Examination', the system locks your paper and starts the countdown.'"
    )

    # =========================================================================
    # SLIDE 5: STAGE 3 - THE ACTIVE EXAMINATION INTERFACE
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header_banner(s5, "STAGE 3 OF 7", "The Active Examination Screen & Top Header", "Stage 3A")
    add_footer(s5)

    s3_img = os.path.join(screenshot_dir, "stage3_exam_screen.png")
    if os.path.exists(s3_img):
        s5.shapes.add_picture(s3_img, Inches(0.6), Inches(1.4), width=Inches(7.2))

    panel = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.4), Inches(4.733), Inches(5.45))
    panel.fill.solid()
    panel.fill.fore_color.rgb = C_WHITE
    panel.line.color.rgb = C_BORDER
    panel.line.width = Pt(1.5)

    ptf = panel.text_frame
    ptf.word_wrap = True
    ptf.margin_left = ptf.margin_right = Inches(0.25)
    ptf.margin_top = Inches(0.25)

    p = ptf.paragraphs[0]
    p.text = "🖥️ EXAM SCREEN ELEMENTS"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY_DARK

    e_steps = [
        ("Candidate Profile Header", "Displays your name, PSN, paper code, and profile initials so you always know you are in your authorized session."),
        ("Live Countdown Timer", "Shows your remaining time in minutes and seconds (e.g. 14:32). Changes to amber at 5 minutes and red at 1 minute."),
        ("Proctor Protected Shield", "Indicates the background anti-cheating monitor is actively securing your test."),
        ("Question Display Card", "Clear, high-contrast question text with single-choice radio options (A, B, C, D)."),
        ("Instant Auto-Save", "Every time you click an option, it is instantly recorded in browser memory and local cache.")
    ]

    for title, desc in e_steps:
        p = ptf.add_paragraph()
        p.text = f"• {title}: "
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_DARK
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_MUTED

    p = ptf.add_paragraph()
    p.text = "\n⏱️ TIMER BEHAVIOR:"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_GOLD

    p = ptf.add_paragraph()
    p.text = "When time reaches 00:00, the system automatically finalizes and submits all chosen answers. You will never lose any answer you selected."
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_DARK

    set_speaker_notes(s5,
        "FACILITATOR SCRIPT:\n"
        "'Now examine Slide 5: this is your examination workspace.\n"
        "At the top, keep your eye on the timer. It counts down continuously from 20:00.\n"
        "You do not need to press any 'Save' button. As soon as you click an option, the system marks it as answered.\n"
        "Even if your network disconnects for a second, your selections are saved securely in your browser cache.'"
    )

    # =========================================================================
    # SLIDE 6: STAGE 3 (CONT.) - 40-QUESTION PALETTE GRID & CONTROLS
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header_banner(s6, "STAGE 3 (CONT.)", "Question Grid (40) & Navigation Controls", "Stage 3B")
    add_footer(s6)

    # Re-use s3_img or zoomed in crop
    if os.path.exists(s3_img):
        s6.shapes.add_picture(s3_img, Inches(0.6), Inches(1.4), width=Inches(7.2))

    panel = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.4), Inches(4.733), Inches(5.45))
    panel.fill.solid()
    panel.fill.fore_color.rgb = C_WHITE
    panel.line.color.rgb = C_BORDER
    panel.line.width = Pt(1.5)

    ptf = panel.text_frame
    ptf.word_wrap = True
    ptf.margin_left = ptf.margin_right = Inches(0.25)
    ptf.margin_top = Inches(0.22)

    p = ptf.paragraphs[0]
    p.text = "🎯 PALETTE COLOR CODES & BUTTONS"
    p.font.name = "Arial"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY_DARK

    p_colors = [
        ("🟩 GREEN SQUARE", "Answered Question", "Indicates that you have successfully selected an option for that question number."),
        ("🟪 PURPLE SQUARE", "Flagged for Review", "Bookmarked so you can quickly return to it before submitting."),
        ("⬜ GRAY SQUARE", "Unanswered Question", "Indicates questions you have not yet answered. You must aim to turn all 40 squares green!"),
        ("🔘 JUMP TO QUESTION", "Direct Grid Navigation", "Click ANY number (1 to 40) on the grid to jump directly to that question instantly."),
        ("✖ CLEAR CHOICE", "Remove Selection", "Allows you to deselect your answer if you want to leave it unselected."),
        ("⚑ FLAG FOR REVIEW", "Mark Difficult Questions", "Bookmarks the question so you can move forward and return later.")
    ]

    for title, sub, desc in p_colors:
        p = ptf.add_paragraph()
        p.text = f"{title} ({sub}): "
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = C_DARK
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_MUTED

    set_speaker_notes(s6,
        "FACILITATOR SCRIPT:\n"
        "'Look closely at the right-hand side of your screen: the Question Grid with 40 boxes.\n"
        "Remember these colors:\n"
        "- GREEN means you have answered.\n"
        "- PURPLE means you flagged it for review.\n"
        "- GRAY means you have NOT answered yet.\n"
        "You do not need to click 'Next' 30 times to get to question 30. Simply tap 30 on the grid!\n"
        "If you change your mind and want to clear an answer, click 'Clear Choice'.'"
    )

    # =========================================================================
    # SLIDE 7: STAGE 4 - PROCTORED BLACKOUT CURTAIN
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header_banner(s7, "STAGE 4 OF 7", "Anti-Cheating & Proctored Blackout Curtain", "Stage 4")
    add_footer(s7)

    s4_img = os.path.join(screenshot_dir, "stage4_blackout_curtain.png")
    if os.path.exists(s4_img):
        s7.shapes.add_picture(s4_img, Inches(0.6), Inches(1.4), width=Inches(7.2))

    panel = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.4), Inches(4.733), Inches(5.45))
    panel.fill.solid()
    panel.fill.fore_color.rgb = C_WHITE
    panel.line.color.rgb = C_BORDER
    panel.line.width = Pt(1.5)

    ptf = panel.text_frame
    ptf.word_wrap = True
    ptf.margin_left = ptf.margin_right = Inches(0.25)
    ptf.margin_top = Inches(0.25)

    p = ptf.paragraphs[0]
    p.text = "🛡️ THE BLACKOUT CURTAIN"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_RED

    b_triggers = [
        ("What Triggers the Curtain?", "Switching to another browser tab, opening another app, minimizing the browser window, or answering phone calls."),
        ("What Happens Instantly?", "The exam interface is immediately blacked out. Questions are hidden to prevent external viewing or search."),
        ("How to Recover?", "Click anywhere on the black curtain to refocus the exam window immediately."),
        ("Consequence of Curtain Activation?", "Refocusing will trigger the official Security Violation Warning Modal.")
    ]

    for title, desc in b_triggers:
        p = ptf.add_paragraph()
        p.text = f"• {title}: "
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_DARK
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_MUTED

    p = ptf.add_paragraph()
    p.text = "\n🚫 STRICTLY PROHIBITED SHORTCUTS:"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_RED

    p = ptf.add_paragraph()
    p.text = "PrintScreen, Win+Shift+S, Cmd+Shift+3/4, Ctrl+P (Print), Ctrl+S (Save), F12 (DevTools), right-click, and text copying are locked and permanently logged."
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_DARK

    set_speaker_notes(s7,
        "FACILITATOR SCRIPT:\n"
        "'Candidates, this is a crucial security rule:\n"
        "If you minimize your browser, switch to WhatsApp, open Google, or try to take a screenshot, "
        "your screen will immediately go pitch black as shown on this slide.\n"
        "The system considers this an unauthorized window blur.\n"
        "Clicking back into the screen will resume the exam, but it will register an infraction against you.\n"
        "Keep your focus solely on the CBT window throughout the 20 minutes.'"
    )

    # =========================================================================
    # SLIDE 8: STAGE 5 - THE 3-STRIKE VIOLATION POLICY
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header_banner(s8, "STAGE 5 OF 7", "Security Violation Warnings (3-Strike System)", "Stage 5")
    add_footer(s8)

    s5_img = os.path.join(screenshot_dir, "stage5_strike_warning.png")
    if os.path.exists(s5_img):
        s8.shapes.add_picture(s5_img, Inches(0.6), Inches(1.4), width=Inches(7.2))

    panel = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.4), Inches(4.733), Inches(5.45))
    panel.fill.solid()
    panel.fill.fore_color.rgb = C_WHITE
    panel.line.color.rgb = C_BORDER
    panel.line.width = Pt(1.5)

    ptf = panel.text_frame
    ptf.word_wrap = True
    ptf.margin_left = ptf.margin_right = Inches(0.25)
    ptf.margin_top = Inches(0.25)

    p = ptf.paragraphs[0]
    p.text = "🚨 THE 3-STRIKE PROGRESSION"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_RED

    strikes_desc = [
        ("STRIKE 1: Formal Warning", "The warning modal pops up. Your incident is recorded with a timestamp on the server. You must click 'Return to Examination'."),
        ("STRIKE 2: Final Warning", "A high-urgency red alert appears: 'You have only ONE warning remaining!'. Any further infraction will terminate the exam."),
        ("STRIKE 3: IMMEDIATE TERMINATION", "Your test session ends instantly! The system auto-submits your answers and flags your file as 'DISQUALIFIED / TERMINATED'.")
    ]

    for title, desc in strikes_desc:
        p = ptf.add_paragraph()
        p.text = f"• {title}: "
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_DARK
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_MUTED

    p = ptf.add_paragraph()
    p.text = "\n🛡️ WATERMARK AUDIT:"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY_DARK

    p = ptf.add_paragraph()
    p.text = "A forensic watermark with your PSN and timestamp is overlaid across the exam page. Attempting to photograph your screen with a second phone will be traced directly to your PSN."
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_DARK

    set_speaker_notes(s8,
        "FACILITATOR SCRIPT:\n"
        "'The Civil Service Commission has implemented a strict 3-Strike policy:\n"
        "- Strike 1: You get an official warning pop-up.\n"
        "- Strike 2: You get a final warning in bold red.\n"
        "- Strike 3: The system terminates your test immediately, locks your answers, and submits your paper with a proctor violation flag.\n"
        "Do not risk your promotion. Stay on the examination page until you finish.'"
    )

    # =========================================================================
    # SLIDE 9: STAGE 6 - UNANSWERED QUESTIONS CHECK
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header_banner(s9, "STAGE 6 OF 7", "Incomplete Examination Check Before Submitting", "Stage 6")
    add_footer(s9)

    s6_img = os.path.join(screenshot_dir, "stage6_unanswered_modal.png")
    if os.path.exists(s6_img):
        s9.shapes.add_picture(s6_img, Inches(0.6), Inches(1.4), width=Inches(7.2))

    panel = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.4), Inches(4.733), Inches(5.45))
    panel.fill.solid()
    panel.fill.fore_color.rgb = C_WHITE
    panel.line.color.rgb = C_BORDER
    panel.line.width = Pt(1.5)

    ptf = panel.text_frame
    ptf.word_wrap = True
    ptf.margin_left = ptf.margin_right = Inches(0.25)
    ptf.margin_top = Inches(0.25)

    p = ptf.paragraphs[0]
    p.text = "⚠️ ANSWER ALL 40 QUESTIONS"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_GOLD

    u_points = [
        ("Incomplete Submission Modal", "If you click 'Submit' while any question is still gray (unanswered), this dialog warns you immediately."),
        ("No Negative Marking", "There is no penalty for guessing. An unanswered question receives zero marks, whereas an educated attempt could earn you a full mark."),
        ("Check the Palette", "Look for any remaining gray squares. Tap on each gray square to answer it."),
        ("Recommended Action", "Click 'Return to Exam & Answer All Questions' to maximize your score potential before your 20 minutes expire.")
    ]

    for title, desc in u_points:
        p = ptf.add_paragraph()
        p.text = f"• {title}: "
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_DARK
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_MUTED

    p = ptf.add_paragraph()
    p.text = "\n💡 GOLDEN RULE:"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_GREEN

    p = ptf.add_paragraph()
    p.text = "Never leave any question blank! Turn all 40 boxes GREEN before clicking Submit."
    p.font.name = "Arial"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY_DARK

    set_speaker_notes(s9,
        "FACILITATOR SCRIPT:\n"
        "'Before submitting, the computer performs an automated safety check.\n"
        "If you have unanswered questions, it will stop you with this warning showing exactly how many you missed.\n"
        "There is no negative marking in Kwara Civil Service CBT exams! "
        "A blank question is a guaranteed zero. An attempted question gives you a chance of scoring. "
        "Always click 'Return to Exam' and attempt all 40 questions.'"
    )

    # =========================================================================
    # SLIDE 10: STAGE 7 - SUBMISSION ACKNOWLEDGEMENT SLIP
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header_banner(s10, "STAGE 7 OF 7", "Official Submission Acknowledgement Slip", "Stage 7")
    add_footer(s10)

    s7_img = os.path.join(screenshot_dir, "stage7_submission_slip.png")
    if os.path.exists(s7_img):
        s10.shapes.add_picture(s7_img, Inches(0.6), Inches(1.4), width=Inches(7.2))

    panel = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.4), Inches(4.733), Inches(5.45))
    panel.fill.solid()
    panel.fill.fore_color.rgb = C_WHITE
    panel.line.color.rgb = C_BORDER
    panel.line.width = Pt(1.5)

    ptf = panel.text_frame
    ptf.word_wrap = True
    ptf.margin_left = ptf.margin_right = Inches(0.25)
    ptf.margin_top = Inches(0.25)

    p = ptf.paragraphs[0]
    p.text = "📄 WHAT YOUR SLIP SHOWS"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY_DARK

    s_items = [
        ("Recorded & Locked Status", "Confirms that your 40 answers have been securely received in the Commission database."),
        ("Candidate Audit Details", "Displays Officer Name, PSN, MDA, Grade Level, Paper Code, Time Spent, and exact Submission Date & Time."),
        ("Electronic Reference Code", "A unique tracking code (e.g. Ref: KWS-CSC-SUB-01048-123456) verifying your submission integrity."),
        ("Print Acknowledgement Slip", "Click the green 'Print Acknowledgement Slip (PDF)' button to save or print a copy for your records."),
        ("Session Concluded", "Your access code is deactivated. Retakes or additional attempts are strictly prohibited.")
    ]

    for title, desc in s_items:
        p = ptf.add_paragraph()
        p.text = f"✔ {title}: "
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_DARK
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_MUTED

    p = ptf.add_paragraph()
    p.text = "\n🔒 ZERO SCORES DISPLAYED:"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY_DARK

    p = ptf.add_paragraph()
    p.text = "Notice that no numerical score or percentage appears on this slip. This is by design per Commission policy."
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_MUTED

    set_speaker_notes(s10,
        "FACILITATOR SCRIPT:\n"
        "'When you click final submit, this is what appears on your screen: your Submission Acknowledgement Slip.\n"
        "It provides official proof that your test was successfully recorded.\n"
        "Check your electronic reference code at the bottom.\n"
        "You can click 'Print Acknowledgement Slip' to save it as a PDF or print it.\n"
        "Once you see this slip, your test is safely completed. You may close your browser and exit quietly.'"
    )

    # =========================================================================
    # SLIDE 11: RESULT CONFIDENTIALITY & MDA COMMUNICATION POLICY
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header_banner(s11, "COMMISSION POLICY", "Result Confidentiality & Official MDA Processing", "Policy")
    add_footer(s11)

    # 3 Policy Columns
    policy_cols = [
        ("🔒", "WORKSTATION PRIVACY", [
            "Examination scores are NOT displayed to candidates on the computer or mobile screen upon completion.",
            "This eliminates hall distractions, emotional reactions, and peer pressure during active batches.",
            "Ensures equal confidentiality across all candidate cadres and ministries."
        ]),
        ("🏛️", "MDA CHANNEL DISPATCH", [
            "All candidate responses and telemetry logs are securely encrypted and transferred to the CSC central server.",
            "Master scores are compiled by the Civil Service Commission alongside seniority, APER scores, and interview results.",
            "Official promotion results are communicated formally through your parent MDA's Director of Personnel Management (DPM)."
        ]),
        ("⚖️", "INTEGRITY & AUDIT TRAIL", [
            "The Admin Portal allows authorized CSC Commissioners to audit test timestamps, time spent, and proctoring telemetry.",
            "Any candidate with logged security strikes is reviewed by the examination malpractice committee.",
            "Complete data protection guarantees that no unofficial result leaks occur."
        ])
    ]

    for idx, (icon, col_title, points) in enumerate(policy_cols):
        x = Inches(0.6 + idx * 4.1)
        card = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(3.9), Inches(5.3))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1.5)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.25)
        ctf.margin_top = Inches(0.3)

        p = ctf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = icon
        p.font.size = Pt(36)

        p = ctf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.text = col_title
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_PRIMARY_DARK

        p = ctf.add_paragraph()
        p.text = ""

        for pt in points:
            p = ctf.add_paragraph()
            p.text = f"• {pt}"
            p.font.name = "Arial"
            p.font.size = Pt(10)
            p.font.color.rgb = C_DARK
            p.space_after = Pt(8)

    set_speaker_notes(s11,
        "FACILITATOR SCRIPT:\n"
        "'We want to emphasize this policy very clearly:\n"
        "Many candidates ask: 'Why didn't I see my score when I finished?'\n"
        "Civil Service promotion is a holistic evaluation. The CBT score is combined with your Annual Performance Evaluation Report (APER), "
        "your seniority, and promotional guidelines.\n"
        "Scores are not displayed on workstations to prevent unrest, hall disorder, or premature conclusions.\n"
        "Your official results will reach you through your Ministry, Department, or Agency through the proper administrative channels.'"
    )

    # =========================================================================
    # SLIDE 12: SMARTPHONE & MOBILE DEVICE BEST PRACTICES
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header_banner(s12, "TECHNICAL ADVISORY", "Best Practices for Mobile Phone / Tablet Candidates", "Mobile Guide")
    add_footer(s12)

    mob_cards = [
        ("🔋 BATTERY PREPARATION", "Charge your mobile device to at least 80% before reporting to the CBT center. Avoid low-power mode dimming during your test.", C_PRIMARY_DARK),
        ("🔕 DO NOT DISTURB (DND)", "Enable 'Do Not Disturb' or airplane mode with Wi-Fi only. Incoming phone calls minimize the browser and trigger security strikes!", C_RED),
        ("📱 WIDE / LANDSCAPE MODE", "Rotate your phone to landscape mode if you prefer seeing both the question card and the 40-question grid side-by-side.", C_PRIMARY_TEAL),
        ("👆 NO 3-FINGER SWIPES", "Do not perform 3-finger swipe gestures on your screen. The proctoring system intercepts this as a screenshot attempt and blocks your screen.", C_GOLD),
        ("🛡️ SCREEN WAKE-LOCK ACTIVE", "The CBT portal automatically keeps your phone screen awake so your device will not sleep during the 20 minutes.", C_PRIMARY_DARK),
        ("🌐 STABLE BROWSER", "Use modern Google Chrome, Microsoft Edge, or Safari. Avoid third-party mini browsers (like Opera Mini) that compress layout.", C_PRIMARY_TEAL)
    ]

    for idx, (title, text, color) in enumerate(mob_cards):
        row = idx // 2
        col = idx % 2
        x = Inches(0.6 + col * 6.2)
        y = Inches(1.5 + row * 1.75)

        card = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.95), Inches(1.55))
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.2)
        ctf.margin_top = Inches(0.18)

        p = ctf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = color

        p = ctf.add_paragraph()
        p.text = text
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK

    set_speaker_notes(s12,
        "FACILITATOR SCRIPT:\n"
        "'For candidates taking the CBT on mobile devices or tablets, these 6 technical tips are essential:\n"
        "Number one: Put your phone on 'Do Not Disturb'. If your relative or colleague calls you during the exam, "
        "the incoming call overlay minimizes the CBT page, and the system registers a Strike 1 violation!\n"
        "Number two: Do not attempt 3-finger screenshot gestures. It will trigger a screen blackout.\n"
        "Number three: Keep your screen clean and use standard Google Chrome or Edge.'"
    )

    # =========================================================================
    # SLIDE 13: GOLDEN CHECKLIST & SUMMARY
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header_banner(s13, "EXAM DAY CHECKLIST", "Golden Checklist & Final Instructions for Candidates", "Checklist")
    add_footer(s13)

    # Two Main Columns: WHAT TO DO (Green) vs WHAT NEVER TO DO (Red)
    col_do = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.5), Inches(5.9), Inches(4.5))
    col_do.fill.solid()
    col_do.fill.fore_color.rgb = RGBColor(240, 253, 244) # Soft green
    col_do.line.color.rgb = C_GREEN
    col_do.line.width = Pt(2)

    dtf = col_do.text_frame
    dtf.word_wrap = True
    dtf.margin_left = dtf.margin_right = Inches(0.25)
    dtf.margin_top = Inches(0.25)

    p = dtf.paragraphs[0]
    p.text = "✅ WHAT YOU MUST DO"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(6, 95, 70)

    dos = [
        "Bring your printed Photocard Slip with your PSN & Code 1.",
        "Arrive at your allocated CBT Hall 30 minutes before your batch time.",
        "Carefully verify your Name, MDA, and Paper Code before clicking Begin.",
        "Keep your eye on the 20-minute countdown clock.",
        "Attempt all 40 questions — ensure all boxes turn GREEN.",
        "Print or save your Submission Acknowledgement Slip upon completion."
    ]
    for d in dos:
        p = dtf.add_paragraph()
        p.text = f"✔  {d}"
        p.font.name = "Arial"
        p.font.size = Pt(10.5)
        p.font.color.rgb = RGBColor(6, 95, 70)
        p.space_after = Pt(6)

    col_dont = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.9), Inches(4.5))
    col_dont.fill.solid()
    col_dont.fill.fore_color.rgb = RGBColor(254, 242, 242) # Soft red
    col_dont.line.color.rgb = C_RED
    col_dont.line.width = Pt(2)

    ntf = col_dont.text_frame
    ntf.word_wrap = True
    ntf.margin_left = ntf.margin_right = Inches(0.25)
    ntf.margin_top = Inches(0.25)

    p = ntf.paragraphs[0]
    p.text = "❌ WHAT YOU MUST NEVER DO"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(153, 27, 27)

    donts = [
        "NEVER switch to another browser tab or application (Strike trigger).",
        "NEVER attempt screen capture (PrintScreen, Snipping tool, shortcuts).",
        "NEVER leave questions unanswered (each blank question is 0 marks).",
        "NEVER click 'Submit' prematurely without double-checking your grid.",
        "NEVER expect or demand a numerical score on your workstation.",
        "NEVER attempt a second login after submitting (Single attempt rule)."
    ]
    for nd in donts:
        p = ntf.add_paragraph()
        p.text = f"✖  {nd}"
        p.font.name = "Arial"
        p.font.size = Pt(10.5)
        p.font.color.rgb = RGBColor(153, 27, 27)
        p.space_after = Pt(6)

    # Bottom encouragement banner
    bot_card = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.15), Inches(12.1), Inches(0.8))
    bot_card.fill.solid()
    bot_card.fill.fore_color.rgb = C_PRIMARY_DARK
    bot_card.line.color.rgb = C_GOLD
    btf = bot_card.text_frame
    bp = btf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    bp.text = "🌟 The Kwara State Civil Service Commission wishes all promotion candidates success in their evaluation!"
    bp.font.name = "Arial"
    bp.font.size = Pt(12)
    bp.font.bold = True
    bp.font.color.rgb = C_WHITE

    set_speaker_notes(s13,
        "FACILITATOR SCRIPT:\n"
        "'To conclude our briefing today:\n"
        "Remember your golden checklist on Slide 13.\n"
        "Be calm, read each question carefully, manage your 20 minutes, answer all 40 questions, "
        "and print your acknowledgement slip when you submit.\n"
        "Are there any questions before we take candidates to the practice demo?'"
    )

    out_file = os.path.abspath("Kwara_CSC_2026_CBT_Candidate_Orientation_Guide.pptx")
    prs.save(out_file)
    print(f"PowerPoint Presentation successfully created: {out_file} (Size: {os.path.getsize(out_file)} bytes)")

if __name__ == "__main__":
    create_presentation()
