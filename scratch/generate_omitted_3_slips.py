import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

omitted_8 = [
    {
        "sn": 1,
        "name": "Oniremu Oluseye Oyetokunbo",
        "psn": "127610",
        "csc_file": "S.16282087",
        "sex": "F",
        "present_rank": "Assistant Chief Nursing Officer",
        "present_gl": "13",
        "proposed_rank": "Chief Nursing Officer",
        "proposed_gl": "14",
        "mda": "KWSUTH",
        "group_category": "GROUP A",
        "exam_code": "KWSUTH/A4",
        "exam_date": "Tuesday, 29th September 2026",
        "batch_session": "Session 1",
        "batch_time": "10:00 AM - 11:00 AM",
        "accreditation_time": "09:30 AM",
        "code_1": "A-24792",
        "phone": "08038271918",
        "email": "127610@cbt.kw.gov.ng",
        "lga": "Ifelodun"
    },
    {
        "sn": 2,
        "name": "Agboola Oyetunji Adeyemi",
        "psn": "135733",
        "csc_file": "S.319152302",
        "sex": "M",
        "present_rank": "Principal  Store Officer II",
        "present_gl": "10",
        "proposed_rank": "Principal  Store Officer I",
        "proposed_gl": "12",
        "mda": "OHOS",
        "group_category": "GROUP B",
        "exam_code": "OHOS/B4",
        "exam_date": "Tuesday, 29th September 2026",
        "batch_session": "Session 5",
        "batch_time": "02:00 PM - 03:00 PM",
        "accreditation_time": "01:30 PM",
        "code_1": "B-97785",
        "phone": "07062739105",
        "email": "135733@cbt.kw.gov.ng",
        "lga": "Oyun"
    },
    {
        "sn": 3,
        "name": "Hammed Ibrahim Olawale",
        "psn": "141443",
        "csc_file": "S.082131",
        "sex": "M",
        "present_rank": "Higher Store Officer",
        "present_gl": "8",
        "proposed_rank": "Senior Store Officer",
        "proposed_gl": "9",
        "mda": "OHOS",
        "group_category": "GROUP C",
        "exam_code": "OHOS/C3",
        "exam_date": "Wednesday, 30th September 2026",
        "batch_session": "Session 2",
        "batch_time": "11:00 AM - 12:00 PM",
        "accreditation_time": "10:30 AM",
        "code_1": "C-61121",
        "phone": "08140505550",
        "email": "141443@cbt.kw.gov.ng",
        "lga": "Offa"
    },
    {
        "sn": 4,
        "name": "Amuzat Aminat",
        "psn": "128645",
        "csc_file": "S.318694",
        "sex": "F",
        "present_rank": "Principal Nursing Officer",
        "present_gl": "12",
        "proposed_rank": "Assistant Chief Nursing Officer",
        "proposed_gl": "13",
        "mda": "KWSUTH",
        "group_category": "GROUP B",
        "exam_code": "KWSUTH/B4",
        "exam_date": "Tuesday, 29th September 2026",
        "batch_session": "Session 3",
        "batch_time": "12:00 PM - 01:00 PM",
        "accreditation_time": "11:30 AM",
        "code_1": "B-30883",
        "phone": "08135545164",
        "email": "128645@cbt.kw.gov.ng",
        "lga": "Ogbomosho North"
    },
    {
        "sn": 5,
        "name": "Akinrinmade Ayoola",
        "psn": "128042",
        "csc_file": "S.318619",
        "sex": "F",
        "present_rank": "Assistant Chief Nursing Supt.",
        "present_gl": "13",
        "proposed_rank": "Chief Nursing Supt.",
        "proposed_gl": "14",
        "mda": "HMB",
        "group_category": "GROUP A",
        "exam_code": "HMB/A3",
        "exam_date": "Tuesday, 29th September 2026",
        "batch_session": "Session 1",
        "batch_time": "10:00 AM - 11:00 AM",
        "accreditation_time": "09:30 AM",
        "code_1": "A-98542",
        "phone": "08023309439",
        "email": "128042@cbt.kw.gov.ng",
        "lga": "Ifelodun"
    },
    {
        "sn": 6,
        "name": "Mohammed Suleiman",
        "psn": "137515",
        "csc_file": "S.708",
        "sex": "M",
        "present_rank": "Senior Nursing Supt.",
        "present_gl": "9",
        "proposed_rank": "Principal Nursing Supt. II",
        "proposed_gl": "10",
        "mda": "HMB",
        "group_category": "GROUP C",
        "exam_code": "HMB/C3",
        "exam_date": "Wednesday, 30th September 2026",
        "batch_session": "Session 1",
        "batch_time": "10:00 AM - 11:00 AM",
        "accreditation_time": "09:30 AM",
        "code_1": "C-21835",
        "phone": "080664938",
        "email": "137515@cbt.kw.gov.ng",
        "lga": "Edu"
    },
    {
        "sn": 7,
        "name": "Giwa Aminat",
        "psn": "135667",
        "csc_file": "135667",
        "sex": "F",
        "present_rank": "CHEW",
        "present_gl": "7",
        "proposed_rank": "Higher (CHEW)",
        "proposed_gl": "8",
        "mda": "PHCDA",
        "group_category": "GROUP D",
        "exam_code": "PHCDA/D1",
        "exam_date": "Wednesday, 30th September 2026",
        "batch_session": "Session 3",
        "batch_time": "12:00 PM - 01:00 PM",
        "accreditation_time": "11:30 AM",
        "code_1": "D-73857",
        "phone": "07066721337",
        "email": "135667@cbt.kw.gov.ng",
        "lga": "Ilorin East"
    },
    {
        "sn": 8,
        "name": "Zubair Dupe Olayinka",
        "psn": "135073",
        "csc_file": "135073",
        "sex": "F",
        "present_rank": "Higher (CHEW)",
        "present_gl": "8",
        "proposed_rank": "Senior (CHEW)",
        "proposed_gl": "9",
        "mda": "PHCDA",
        "group_category": "GROUP C",
        "exam_code": "PHCDA/C1",
        "exam_date": "Wednesday, 30th September 2026",
        "batch_session": "Session 3",
        "batch_time": "12:00 PM - 01:00 PM",
        "accreditation_time": "11:30 AM",
        "code_1": "C-48205",
        "phone": "08039401871",
        "email": "135073@cbt.kw.gov.ng",
        "lga": "Ifelodun"
    }
]

# Generate standalone Excel
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Omitted Candidates Batch 3"

header_fill = PatternFill(start_color="004D40", end_color="004D40", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)
font_bold = Font(name="Calibri", size=11, bold=True)
font_regular = Font(name="Calibri", size=11)
code_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
code_font = Font(name="Consolas", size=12, bold=True, color="B45309")

headers = [
    "S/N", "PSN", "Candidate Full Name", "Registration Clearance Code (Code 1)",
    "Proposed Rank", "Grade Level", "MDA", "Group Category", "Exam Code",
    "Exam Date", "Batch Session", "Batch Time", "Accreditation Time", "Onboarding Steps"
]
ws.append(headers)
for col_idx in range(1, len(headers) + 1):
    c = ws.cell(row=1, column=col_idx)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")

for idx, c in enumerate(omitted_8, start=1):
    row_vals = [
        idx,
        c["psn"],
        c["name"],
        c["code_1"],
        c["proposed_rank"],
        c["proposed_gl"],
        c["mda"],
        c["group_category"],
        c["exam_code"],
        c["exam_date"],
        c["batch_session"],
        c["batch_time"],
        c["accreditation_time"],
        "1. Enter PSN + Code 1 -> 2. Upload Passport -> 3. Print Photocard"
    ]
    ws.append(row_vals)
    r_idx = idx + 1
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=r_idx, column=col_idx)
        cell.border = border_thin
        if col_idx == 4:
            cell.fill = code_fill
            cell.font = code_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in [1, 2, 6, 8, 9, 10, 11, 12, 13]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = font_bold if col_idx in [2, 6] else font_regular
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.font = font_bold if col_idx == 3 else font_regular

for col in ws.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

excel_name = "Kwara_CSC_2026_CBT_Omitted_Candidates_Batch_3_(8_Officers).xlsx"
wb.save(excel_name)
print(f"Generated {excel_name}")

# Generate HTML Slips
html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>Kwara CSC 2026 CBT - Registration Clearance Slips (Omitted Candidates Batch 3 - 8 Officers)</title>
  <style>
    *, *::before, *::after {{
      box-sizing: border-box;
      margin: 0;
      padding: 0;
    }}
    @page {{
      size: A4 portrait;
      margin: 6mm 6mm;
    }}
    body {{
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif;
      background: #f1f5f9;
      color: #0f172a;
      -webkit-print-color-adjust: exact !important;
      print-color-adjust: exact !important;
    }}
    .no-print {{
      background: #004d40;
      color: #fff;
      padding: 14px 20px;
      display: flex;
      flex-wrap: wrap;
      align-items: center;
      justify-content: space-between;
      gap: 12px;
      position: sticky;
      top: 0;
      z-index: 1000;
      box-shadow: 0 4px 12px rgba(0,0,0,0.18);
    }}
    .no-print h1 {{
      font-size: 1.1rem;
      font-weight: 800;
    }}
    .no-print p {{
      font-size: 0.8rem;
      opacity: 0.9;
    }}
    .btn-print {{
      background: #f59e0b;
      color: #000;
      font-weight: 800;
      border: none;
      padding: 9px 20px;
      border-radius: 6px;
      cursor: pointer;
      font-size: 0.92rem;
    }}
    .btn-print:hover {{
      background: #d97706;
      color: #fff;
    }}
    .screen-wrapper {{
      max-width: 210mm;
      margin: 15px auto;
    }}
    .page-container {{
      background: #fff;
      width: 198mm;
      min-height: 284mm;
      margin: 0 auto 20px auto;
      padding: 4mm 4mm;
      box-shadow: 0 4px 15px rgba(0,0,0,0.08);
      display: grid;
      grid-template-columns: 1fr 1fr;
      grid-template-rows: repeat(4, 1fr);
      gap: 3.5mm 3.5mm;
      page-break-after: always;
      break-after: page;
    }}
    .slip-card {{
      border: 1.2px dashed #004d40;
      border-radius: 4px;
      padding: 4px 6px;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
      background: #fff;
    }}
    .slip-header {{
      border-bottom: 1.2px solid #004d40;
      padding-bottom: 2px;
      margin-bottom: 3px;
      text-align: center;
    }}
    .header-seal {{
      font-size: 7.8px;
      font-weight: 900;
      color: #004d40;
      letter-spacing: 0.02em;
    }}
    .header-sub {{
      font-size: 6.2px;
      font-weight: 700;
      color: #475569;
    }}
    .code-box {{
      background: #fffbeb;
      border: 1.4px dashed #d97706;
      border-radius: 3px;
      padding: 2.5px 4px;
      margin-bottom: 3px;
      display: flex;
      justify-content: space-between;
      align-items: center;
    }}
    .code-lbl {{
      font-size: 6.2px;
      font-weight: 700;
      color: #92400e;
    }}
    .code-val {{
      font-family: "Courier New", monospace;
      font-size: 11px;
      font-weight: 900;
      color: #b45309;
      letter-spacing: 0.06em;
    }}
    .details-table {{
      width: 100%;
      border-collapse: collapse;
      font-size: 7px;
      margin-bottom: 3px;
    }}
    .details-table td {{
      padding: 1.2px 2px;
      vertical-align: middle;
      line-height: 1.15;
    }}
    .td-lbl {{
      font-weight: 700;
      color: #475569;
      width: 25%;
    }}
    .td-val {{
      font-weight: 800;
      color: #0f172a;
    }}
    .psn-highlight {{
      color: #004d40;
      font-size: 9px;
      font-weight: 900;
    }}
    .instruction-box {{
      background: #f8fafc;
      border-left: 2.5px solid #004d40;
      padding: 2px 4px;
      font-size: 6.2px;
      color: #334155;
      line-height: 1.2;
      border-radius: 0 2px 2px 0;
      margin-bottom: 2px;
    }}
    .slip-footer {{
      border-top: 0.8px dashed #cbd5e1;
      padding-top: 1.5px;
      display: flex;
      justify-content: space-between;
      font-size: 6px;
      color: #64748b;
    }}
    .slip-footer strong {{
      color: #004d40;
    }}
    @media print {{
      body {{ background: #fff; }}
      .no-print {{ display: none !important; }}
      .screen-wrapper {{ max-width: 100%; margin: 0; padding: 0; }}
      .page-container {{ box-shadow: none; margin: 0; width: 198mm; height: 284mm; }}
    }}
  </style>
</head>
<body>
  <div class="no-print">
    <div>
      <h1>KWARA STATE CIVIL SERVICE COMMISSION</h1>
      <p>2026 Promotion Evaluation CBT &bull; Candidate Registration Clearance Slips (Omitted Batch 3 - 8 Officers)</p>
    </div>
    <div>
      <button class="btn-print" onclick="window.print()">🖨️ Print Slips (PDF / Paper)</button>
    </div>
  </div>
  <div class="screen-wrapper">
    <div class="page-container">
"""

for c in omitted_8:
    html_content += f"""      <div class="slip-card">
        <div class="slip-header">
          <div class="header-seal">KWARA STATE CIVIL SERVICE COMMISSION</div>
          <div class="header-sub">2026 PROMOTION EVALUATION &bull; ONBOARDING CLEARANCE SLIP</div>
        </div>
        <div class="code-box">
          <span class="code-lbl">&#128273; REGISTRATION CODE (CODE 1):</span>
          <span class="code-val">{c['code_1']}</span>
        </div>
        <table class="details-table">
          <tr>
            <td class="td-lbl">OFFICER:</td>
            <td class="td-val">{c['name']}</td>
          </tr>
          <tr>
            <td class="td-lbl">PSN:</td>
            <td class="td-val"><span class="psn-highlight">{c['psn']}</span> &bull; <strong>{c['mda']}</strong></td>
          </tr>
          <tr>
            <td class="td-lbl">RANK / GL:</td>
            <td class="td-val">{c['proposed_rank']} &bull; <strong>GL {c['proposed_gl']}</strong></td>
          </tr>
          <tr>
            <td class="td-lbl">EXAM CODE:</td>
            <td class="td-val"><strong>{c['exam_code']}</strong> ({c['group_category']})</td>
          </tr>
          <tr>
            <td class="td-lbl">DATE / TIME:</td>
            <td class="td-val">{c['exam_date']} &bull; <strong>{c['batch_session']}</strong> ({c['batch_time']})</td>
          </tr>
          <tr>
            <td class="td-lbl">ACCREDITATION:</td>
            <td class="td-val"><strong>{c['accreditation_time']}</strong> &bull; Kwara State CBT Centre</td>
          </tr>
        </table>
        <div class="instruction-box">
          <strong>ONBOARDING INSTRUCTIONS:</strong><br>
          1. Visit the portal, click <em>Candidate Registration</em>, and enter your <strong>PSN ({c['psn']})</strong> and <strong>Code 1 ({c['code_1']})</strong>.<br>
          2. Upload passport photo and print your Official Photocard.<br>
          3. Bring your printed Photocard to the CBT Centre for accreditation.
        </div>
        <div class="slip-footer">
          <span>CSC File: <strong>{c['csc_file']}</strong></span>
          <span>Batch: <strong>Omitted Batch 3 (2026)</strong></span>
          <span>Security Token Required at CBT</span>
        </div>
      </div>
"""

html_content += """    </div>
  </div>
</body>
</html>
"""

# Save HTML files
html_paths = [
    "Omitted_Candidates_Registration_Slips_(8_Candidates).html",
    os.path.join("Kwara_CSC_2026_CBT_Slips_By_MDA", "Omitted_Candidates_Registration_Slips_(8_Candidates).html"),
    os.path.join("kwara_cbt_app", "static", "slips_omitted_8.html"),
    os.path.join("public", "slips_omitted_8.html")
]

for hp in html_paths:
    os.makedirs(os.path.dirname(os.path.abspath(hp)), exist_ok=True)
    with open(hp, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"Generated HTML Slip: {hp}")

print("\nDone generating all slips and Excel workbook!")
